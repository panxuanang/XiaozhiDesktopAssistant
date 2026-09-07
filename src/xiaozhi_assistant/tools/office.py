from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from ..action_log import record
from ..config import load_settings
from .files import resolve_user_path


def _apply_default_word_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)


def create_word_document(path: str, title: str, content: str, author: str = "") -> str:
    target = resolve_user_path(path)
    if target.suffix.lower() != ".docx":
        target = target.with_suffix(".docx")
    target.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _apply_default_word_style(doc)
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title.strip())
    run.bold = True
    run.font.name = "黑体"
    run.font.size = Pt(20)
    for block in [b.strip() for b in content.replace("\r\n", "\n").split("\n")]:
        if not block:
            doc.add_paragraph()
            continue
        p = doc.add_paragraph(block)
        p.paragraph_format.first_line_indent = Pt(24)
        p.paragraph_format.line_spacing = 1.5
    if author:
        p = doc.add_paragraph(author)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.save(target)
    record("create_word_document", {"path": str(target), "title": title}, detail=f"chars={len(content)}")
    return str(target)


def write_material(requirement: str, path: str = "桌面/材料.docx", title: str = "") -> str:
    from ..llm import AIClient
    ai = AIClient()
    prompt = f"""请根据以下要求直接撰写一份可交付的中文材料。\n\n用户要求：{requirement}\n\n要求：结构完整、事实不编造、语言符合用户要求；不要解释写作过程。"""
    content = ai.complete(prompt, system="你是一名专业中文材料写作助手，擅长通知、总结、汇报、方案、会议纪要和商务材料。")
    if not title:
        title = ai.complete(f"给下面材料拟一个简洁正式的标题，只输出标题：\n{content[:4000]}", max_output_tokens=80).strip("《》 \n")
    saved = create_word_document(path, title, content)
    return f"材料已生成并保存：{saved}"


def read_word_document(path: str, max_chars: int = 30000) -> str:
    target = resolve_user_path(path)
    doc = Document(target)
    text = "\n".join(p.text for p in doc.paragraphs)
    return text[:max_chars]


def rewrite_word_document(path: str, instruction: str, output_path: str = "") -> str:
    from ..llm import AIClient
    src = resolve_user_path(path)
    original = read_word_document(str(src))
    ai = AIClient()
    revised = ai.complete(
        f"修改要求：{instruction}\n\n原文：\n{original}",
        system="你是专业文稿编辑。保持原文核心事实，不添加未经提供的事实，直接输出修改后的完整正文。",
    )
    target = resolve_user_path(output_path) if output_path else src.with_name(src.stem + "_修改版.docx")
    saved = create_word_document(str(target), src.stem, revised)
    return f"已生成修改版：{saved}"


def excel_profile(path: str, max_rows: int = 50000) -> dict[str, Any]:
    target = resolve_user_path(path)
    book = pd.ExcelFile(target)
    result: dict[str, Any] = {"file": str(target), "sheets": {}}
    for sheet in book.sheet_names[:20]:
        df = pd.read_excel(target, sheet_name=sheet, nrows=max_rows)
        info: dict[str, Any] = {
            "rows": int(len(df)),
            "columns": [str(c) for c in df.columns],
            "missing": {str(k): int(v) for k, v in df.isna().sum().items() if int(v) > 0},
            "sample": df.head(8).where(pd.notna(df.head(8)), None).to_dict(orient="records"),
        }
        num = df.select_dtypes(include="number")
        if not num.empty:
            desc = num.describe().round(4).to_dict()
            info["numeric_summary"] = desc
            info["numeric_totals"] = {str(k): (None if pd.isna(v) else float(v)) for k, v in num.sum(numeric_only=True).items()}
        cat_summary: dict[str, Any] = {}
        for col in df.select_dtypes(exclude="number").columns[:12]:
            values = df[col].astype(str).value_counts(dropna=True).head(10)
            cat_summary[str(col)] = {str(k): int(v) for k, v in values.items()}
        if cat_summary:
            info["top_values"] = cat_summary
        result["sheets"][sheet] = info
    return result


def analyze_excel(path: str, question: str = "请分析这份表格的关键结论、异常和建议", report_path: str = "") -> str:
    from ..llm import AIClient
    from ..llm.client import json_for_ai
    profile = excel_profile(path)
    settings = load_settings()
    compact = json_for_ai(profile, max_chars=settings.safety.max_file_chars_for_ai)
    ai = AIClient()
    answer = ai.complete(
        f"用户问题：{question}\n\n下面是本地程序对 Excel 做的结构化统计摘要：\n{compact}",
        system="你是资深数据分析师。基于提供的统计结果作答；不要臆造原表中没有的数据；用清晰中文给出结论、证据和可执行建议。",
    )
    record("analyze_excel", {"path": path, "question": question}, detail=answer[:500])
    if report_path:
        saved = create_word_document(report_path, "Excel 数据分析报告", answer)
        return answer + f"\n\n分析报告已保存：{saved}"
    return answer


def excel_preview(path: str, sheet: str = "", rows: int = 20) -> str:
    target = resolve_user_path(path)
    df = pd.read_excel(target, sheet_name=sheet if sheet else 0, nrows=max(1, min(rows, 100)))
    return df.to_csv(index=False)
